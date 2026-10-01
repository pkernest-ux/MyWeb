#!/usr/bin/env python3
"""重建 血液數據 活頁簿中的 平均值／最小值／最大值／術前最近值 工作表。

規則
  * INDEX 日期 = TEST_DATE - FIRST_INTERVAL（每位病人唯一）。
  * INDEX 之前的手術不列入；彙總只納入 INDEX 前 7 天（含 INDEX 當天）及其後的檢驗。
  * 遇到 INDEX 之後第一筆含 UTUC 或 UBUC 的手術日（含當天）即停止；其他類別手術只記錄、不停止。
  * 平均值／最小值／最大值：每位病人 x FIRST_TIME_GROUP 一列，CATEGORY／OP 為 INDEX 手術。
  * 術前最近值：INDEX 起每一筆手術各一列，取手術前 7 天（D6 至 D0，含手術當天）各項目最近值。
  * 平均值四捨五入至小數 2 位；'<1.00'、'>25000.00' 之類截尾值去掉符號後當數值計算，
    其他文字視為缺值。術前最近值則保留原始內容。

用法:  python3 build_summary.py 輸入.xlsx 輸出.xlsx
"""
import statistics
import sys
import warnings

import numpy as np
import openpyxl
import pandas as pd

warnings.filterwarnings('ignore')

GORD = ['術前7天內', '6個月內', '6-12個月', '12-24個月', '24-60個月', '60個月以上']
CATS = ['UTUC', 'UBUC', 'OTHER', 'RCC', 'URETHRA']
KEY = ['IDNO', 'OPDATE', 'CATEGORY', 'OP']


def is_stop_cat(cat):
    return isinstance(cat, str) and ('UTUC' in cat or 'UBUC' in cat)


def same(a, b):
    if pd.isna(a) and pd.isna(b):
        return True
    if pd.isna(a) or pd.isna(b):
        return False
    try:
        return abs(float(a) - float(b)) < 1e-9
    except (TypeError, ValueError):
        return str(a) == str(b)


def load_blood(path):
    raw = pd.read_excel(path, sheet_name='血液數據', dtype=object)
    for c in ['TEST_DATE', 'OPDATE']:
        raw[c] = pd.to_datetime(raw[c])
    for c in ['FIRST_INTERVAL', 'RECENT_INTERVAL', 'TEST_SEQ']:
        raw[c] = pd.to_numeric(raw[c])
    labs = list(raw.columns[8:])
    vals = {c: pd.to_numeric(raw[c].astype(str).str.strip().str.replace(r'^[<>]', '', regex=True),
                             errors='coerce') for c in labs}
    b = pd.concat([raw.iloc[:, :8], pd.DataFrame(vals)], axis=1)
    b['IDX'] = b.TEST_DATE - pd.to_timedelta(b.FIRST_INTERVAL, 'D')
    m = ((b.TEST_DATE.dt.year - b.IDX.dt.year) * 12 + (b.TEST_DATE.dt.month - b.IDX.dt.month)
         - (b.TEST_DATE.dt.day < b.IDX.dt.day).astype(int))
    fi = b.FIRST_INTERVAL
    b['G'] = np.select([fi <= 0, m < 6, m < 12, m < 24, m < 60], GORD[:5], GORD[5])
    return b, raw[labs], labs


def patient_frames(b):
    """每位病人的 INDEX 日期、INDEX 手術、停止日；以及手術清單（來自血液數據的 OPDATE）。"""
    surg = b[KEY].drop_duplicates().sort_values(['IDNO', 'OPDATE']).reset_index(drop=True)
    idx = b.groupby('IDNO').IDX.agg(['min', 'max'])
    assert (idx['min'] == idx['max']).all(), 'INDEX 日期不唯一'
    rows = []
    for pid, idate in idx['min'].items():
        s = surg[surg.IDNO == pid]
        cur = s[s.OPDATE == idate]
        after = s[(s.OPDATE > idate) & s.CATEGORY.map(is_stop_cat)]
        rows.append(dict(IDNO=pid, INDEX_DATE=idate,
                         CATEGORY=cur.CATEGORY.iloc[0] if len(cur) else None,
                         OP=cur.OP.iloc[0] if len(cur) else None,
                         INDEX_IN_LIST=len(cur) > 0,
                         STOP=after.OPDATE.min() if len(after) else pd.NaT))
    return pd.DataFrame(rows), surg


def fmean2(s):
    v = s.dropna()
    return round(statistics.fmean(v), 2) if len(v) else np.nan


def build_summaries(b, labs, pat):
    info = pat.set_index('IDNO')
    stop = b.IDNO.map(info.STOP)
    keep = b[(b.FIRST_INTERVAL >= -6) & (stop.isna() | (b.TEST_DATE < stop))]
    g = keep.groupby(['IDNO', 'G'])[labs]
    out = {}
    for name, df in [('平均值', g.agg(fmean2)), ('最小值', g.min()), ('最大值', g.max())]:
        df = df.reset_index().rename(columns={'G': 'FIRST_TIME_GROUP'})
        df['OPDATE'] = df.IDNO.map(info.INDEX_DATE)
        df['CATEGORY'] = df.IDNO.map(info.CATEGORY)
        df['OP'] = df.IDNO.map(info.OP)
        df['_o'] = df.FIRST_TIME_GROUP.map(GORD.index)
        df = df.sort_values(['IDNO', '_o'])
        out[name] = df[['IDNO', 'OPDATE', 'CATEGORY', 'OP', 'FIRST_TIME_GROUP'] + labs].reset_index(drop=True)
    return out, keep


def build_preop(b, raw, labs, pat, surg):
    info = pat.set_index('IDNO')
    rows = []
    for pid, d in b.groupby('IDNO'):
        d = d.sort_values(['TEST_DATE', 'TEST_SEQ'], na_position='first')
        r_ = raw.loc[d.index]
        for _, s in surg[(surg.IDNO == pid) & (surg.OPDATE >= info.INDEX_DATE[pid])].iterrows():
            sel = (d.TEST_DATE >= s.OPDATE - pd.Timedelta(days=6)) & (d.TEST_DATE <= s.OPDATE)
            if not sel.any():
                continue
            w = r_[sel.values]
            last = {c: (w[c].dropna().iloc[-1] if w[c].notna().any() else np.nan) for c in labs}
            rows.append({'IDNO': pid, 'OPDATE': s.OPDATE, 'CATEGORY': s.CATEGORY, 'OP': s.OP, **last})
    return pd.DataFrame(rows, columns=KEY + labs)


def group_stats(avg):
    rows = []
    for gname in GORD:
        d = avg[avg.FIRST_TIME_GROUP == gname]
        rows.append((gname, '全部', d.IDNO.nunique()))
        for c in CATS:
            rows.append((gname, c, d[d.CATEGORY.fillna('').str.contains(c)].IDNO.nunique()))
    return pd.DataFrame(rows, columns=['FIRST_TIME_GROUP', 'CATEGORY', 'PATIENT_COUNT'])


# ---------------------------------------------------------------- 品管（獨立寫法重算）
def qa_checks(src, b, raw, labs, pat, surg, new, pre, stats):
    info = pat.set_index('IDNO')
    out = []

    # QA06 鍵值唯一
    dup = sum(int(new[k].duplicated(['IDNO', 'FIRST_TIME_GROUP']).sum()) for k in new)
    out.append(('QA06', '三張彙總表鍵值唯一（IDNO×FIRST_TIME_GROUP）', sum(len(v) for v in new.values()), dup))

    # QA07 停止點：逐人以迴圈確認停止日，並確認彙總只納入停止日之前的檢驗
    err = 0
    for pid, d in b.groupby('IDNO'):
        s = surg[surg.IDNO == pid]
        stop = None
        for _, r in s.sort_values('OPDATE').iterrows():
            if r.OPDATE > info.INDEX_DATE[pid] and is_stop_cat(r.CATEGORY):
                stop = r.OPDATE
                break
        if not ((stop is None and pd.isna(info.STOP[pid])) or (stop is not None and stop == info.STOP[pid])):
            err += 1
    out.append(('QA07', '停止日 = INDEX 後第一筆含 UTUC／UBUC 手術日', len(pat), err))

    # QA08 彙總表逐人逐組迴圈重算
    err = 0
    n = 0
    for pid, d in b.groupby('IDNO'):
        st = info.STOP[pid]
        d = d[d.FIRST_INTERVAL >= -6]
        if pd.notna(st):
            d = d[d.TEST_DATE < st]
        for gname in GORD:
            sel = d[d.G == gname]
            for name, fn in [('平均值', lambda v: round(statistics.fmean(v), 2)), ('最小值', min), ('最大值', max)]:
                row = new[name][(new[name].IDNO == pid) & (new[name].FIRST_TIME_GROUP == gname)]
                if sel.empty:
                    err += len(row)
                    continue
                n += 1
                if len(row) != 1:
                    err += 1
                    continue
                for c in labs:
                    v = sel[c].dropna().tolist()
                    if not same(fn(v) if v else np.nan, row[c].iloc[0]):
                        err += 1
                        break
    out.append(('QA08', '平均值／最小值／最大值逐人逐組重算', n, err))

    # QA09 術前最近值逐筆重算（保留原始內容）
    err = 0
    for _, r in pre.iterrows():
        d = b[b.IDNO == r.IDNO]
        sel = d[(d.TEST_DATE >= r.OPDATE - pd.Timedelta(days=6)) & (d.TEST_DATE <= r.OPDATE)]
        sel = sel.sort_values(['TEST_DATE', 'TEST_SEQ'], na_position='first')
        for c in labs:
            v = raw.loc[sel.index, c].dropna()
            if not same(v.iloc[-1] if len(v) else np.nan, r[c]):
                err += 1
                break
    out.append(('QA09', '術前最近值D6至D0逐筆重算（INDEX起每筆手術）', len(pre), err))

    # QA10 分組患者統計
    chk = group_stats(new['平均值']).merge(stats, on=['FIRST_TIME_GROUP', 'CATEGORY'], suffixes=('_a', '_b'))
    out.append(('QA10', 'FIRST_TIME_GROUP患者數與癌症類別統計', len(stats),
                int((chk.PATIENT_COUNT_a != chk.PATIENT_COUNT_b).sum()) + len(stats) - len(chk)))

    # QA11 與原表比對：只有單一手術的病人不受停止規則影響
    O = {k: pd.read_excel(src, sheet_name=k) for k in ['平均值', '最小值', '最大值']}
    one = b.groupby('IDNO').OPDATE.nunique()
    ids = set(one[one == 1].index) & set(pat[pat.INDEX_IN_LIST].IDNO)
    tot = 0
    err = 0
    for k in O:
        o = O[k][O[k].IDNO.isin(ids)].set_index(['IDNO', 'FIRST_TIME_GROUP'])
        n_ = new[k][new[k].IDNO.isin(ids)].set_index(['IDNO', 'FIRST_TIME_GROUP'])
        tot += len(o)
        for key in o.index:
            if key not in n_.index:
                err += 1
                continue
            if not all(same(pd.to_numeric(o.loc[key, c], errors='coerce'), n_.loc[key, c]) for c in labs):
                err += 1
    out.append(('QA11', '與原表比對（僅INDEX在清單內的單一手術病人；差異為平均值恰在0.005的進位）', tot, err))

    # QA12 INDEX 手術不在手術清單
    out.append(('QA12', 'INDEX手術不在血液數據手術清單的病人（CATEGORY／OP留空）', len(pat), int((~pat.INDEX_IN_LIST).sum())))
    return out


# ---------------------------------------------------------------- 輸出
def to_cell(v):
    if v is pd.NaT or (isinstance(v, float) and np.isnan(v)):
        return None
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime()
    if isinstance(v, np.generic):
        return v.item()
    return v


def write_workbook(src, dst, new, pre, stats, qa, notes):
    wb = openpyxl.load_workbook(src)

    def replace(name, df, keep_rows=1):
        ws = wb[name]
        header = [c.value for c in ws[1]]
        assert header == list(df.columns), f'{name} 欄位不一致'
        fmt = {i: ws.cell(2, i + 1).number_format for i in range(len(header))} if ws.max_row >= 2 else {}
        ws.delete_rows(2, ws.max_row)
        for rec in df.itertuples(index=False):
            ws.append([to_cell(v) for v in rec])
        for i, f in fmt.items():
            if f and f != 'General':
                for r in range(2, ws.max_row + 1):
                    ws.cell(r, i + 1).number_format = f

    for k in ['平均值', '最小值', '最大值']:
        replace(k, new[k])
    replace('術前最近值', pre)
    replace('分組患者統計', stats)

    ws = wb['QA_Summary']
    ws.delete_rows(7, ws.max_row)  # 保留 QA01–QA05（血液數據本身未改動）
    for cid, desc, n, err in qa:
        status = 'PASS' if err == 0 else ('CHECK' if cid in ('QA11', 'QA12') else 'FAIL')
        ws.append([cid, desc, n, err, status])

    if '修訂說明' in wb.sheetnames:
        del wb['修訂說明']
    ws = wb.create_sheet('修訂說明')
    for line in notes:
        ws.append([line])
    ws.column_dimensions['A'].width = 130
    wb.save(dst)


def main(src, dst):
    b, raw, labs = load_blood(src)
    pat, surg = patient_frames(b)
    new, keep = build_summaries(b, labs, pat)
    pre = build_preop(b, raw, labs, pat, surg)
    stats = group_stats(new['平均值'])
    qa = qa_checks(src, b, raw, labs, pat, surg, new, pre, stats)
    missing = pat[~pat.INDEX_IN_LIST]
    notes = [
        '修訂說明：平均值／最小值／最大值／術前最近值／分組患者統計 已依下列規則重建（血液數據、CODING BOOK 未改動）。',
        '1. INDEX 日期 = TEST_DATE − FIRST_INTERVAL。INDEX 之前的手術不列入。',
        '2. 平均值／最小值／最大值：每位病人 × FIRST_TIME_GROUP 一列。納入 INDEX 前 7 天（含 INDEX 當天）及其後的檢驗，'
        '遇到 INDEX 之後第一筆含 UTUC 或 UBUC 的手術日（含當天）即停止；其他類別手術只記錄、不停止。',
        '   OPDATE／CATEGORY／OP 為 INDEX 手術的資料。FIRST_INTERVAL = 0（手術當天）歸入「術前7天內」。',
        '3. 術前最近值：INDEX 起每一筆手術各一列，取手術前 7 天（D6 至 D0，含手術當天）各項目最近值（保留原始內容，如「<1.00」）。',
        '4. 數值處理：「<1.00」「>25000.00」去掉符號後當數值計算；其他文字視為缺值。平均值四捨五入至小數 2 位。',
        '5. 手術清單來自血液數據的 OPDATE。若某筆手術之後、下一筆手術之前沒有任何檢驗，該手術不會出現在清單中，'
        '術前最近值也不會有該筆手術。如需完整清單，請提供病理主檔。',
        f'6. 下列 {len(missing)} 位病人的 INDEX 手術不在上述清單，其 CATEGORY／OP 留空：',
    ] + [f'   {r.IDNO}（INDEX 日期 {r.INDEX_DATE.date()}）' for r in missing.itertuples()]
    write_workbook(src, dst, new, pre, stats, qa, notes)
    print('病人', len(pat), '| 彙總列數', {k: len(v) for k, v in new.items()}, '| 術前列數', len(pre))
    for q in qa:
        print(q)


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
